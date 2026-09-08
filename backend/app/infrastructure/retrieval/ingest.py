from __future__ import annotations

import logging
import time
from io import BytesIO
from pathlib import Path

from app.core.ingest_context import ingest_source_key
from app.domain.models import Chunk
from app.infrastructure.llm.embeddings import get_embedder
from app.infrastructure.retrieval.chunking import (
    XLSX_ROWS_PER_CHUNK,
    cohesive_pdf_chunks,
    hybrid_markdown_chunks,
    make_chunk,
    plan_for_suffix,
    recursive_split,
    split_markdown_sections,
)
from app.infrastructure.retrieval.cleaning import clean_cell, clean_prose, looks_like_signal
from app.infrastructure.retrieval.semantic import semantic_pdf_units, semantic_section_units

log = logging.getLogger(__name__)


def _title_from_text(text: str, fallback: str) -> str:
    for line in text.splitlines():
        stripped = line.strip().lstrip("#").strip()
        if stripped and not stripped.startswith("[Page "):
            return stripped[:160]
    return fallback


def _as_of_from_text(text: str) -> str:
    for line in text.splitlines():
        lower = line.lower()
        if "as of" in lower or "effective as of" in lower:
            return line.strip()[:80]
    return "unknown"


def _parse_pdf_pages(data: bytes) -> list[tuple[int, str]]:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    pages: list[tuple[int, str]] = []
    for i, page in enumerate(reader.pages, start=1):
        raw = (page.extract_text() or "").strip()
        cleaned = clean_prose(raw, pdf=True)
        if looks_like_signal(cleaned):
            pages.append((i, cleaned))
    return pages


def _parse_xlsx_sheets(data: bytes) -> list[tuple[str, list[str], list[list[str]]]]:
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    sheets: list[tuple[str, list[str], list[list[str]]]] = []
    try:
        for sheet in wb.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            headers = [clean_cell(c) for c in rows[0]]
            body: list[list[str]] = []
            for row in rows[1:]:
                cells = [clean_cell(c) for c in row]
                if any(cells):
                    body.append(cells)
            if headers or body:
                sheets.append((sheet.title, headers, body))
    finally:
        wb.close()
    return sheets


def ingest_bytes(relative_path: str, data: bytes) -> list[Chunk]:
    """ETL: parse → clean → hybrid chunk. One document in, retrieval units out."""
    token = ingest_source_key.set(relative_path)
    started = time.perf_counter()
    suffix = Path(relative_path).suffix.lower()
    file_id = Path(relative_path).with_suffix("").as_posix()
    fallback_title = Path(relative_path).stem.replace("_", " ")
    plan = plan_for_suffix(suffix)
    pages = 0
    log.info(
        "Chunking start %s (%s, %s bytes, plan=%s)",
        relative_path,
        suffix or "none",
        len(data),
        plan.strategy,
        extra={
            "stage": "chunk_start",
            "suffix": suffix or "none",
            "bytes": len(data),
            "strategy": plan.strategy,
        },
    )
    try:
        if suffix == ".pdf":
            parsed = _parse_pdf_pages(data)
            pages = len(parsed)
            embedder = get_embedder()
            chunks = (
                _chunk_pdf_semantic(file_id, fallback_title, parsed, embedder)
                if embedder
                else _chunk_pdf(file_id, fallback_title, parsed, plan)
            )
        elif suffix in {".xlsx", ".xlsm"}:
            chunks = _chunk_xlsx(file_id, fallback_title, _parse_xlsx_sheets(data), plan)
        elif suffix == ".jsonl":
            text = clean_prose(data.decode("utf-8"))
            chunks = _chunk_jsonl(file_id, fallback_title, text, plan)
        elif suffix == ".md":
            text = clean_prose(data.decode("utf-8"))
            embedder = get_embedder()
            chunks = (
                _chunk_markdown_semantic(file_id, fallback_title, text, embedder)
                if embedder
                else _chunk_markdown(file_id, fallback_title, text, plan)
            )
        else:
            text = clean_prose(data.decode("utf-8"))
            chunks = _chunk_recursive(file_id, fallback_title, text, plan, doc_type="text")
    except Exception:
        log.exception(
            "Chunking failed %s",
            relative_path,
            extra={"stage": "chunk_error", "suffix": suffix or "none"},
        )
        raise
    else:
        sizes = [len(c.text) for c in chunks]
        strategy = chunks[0].strategy if chunks else plan.strategy
        elapsed = int((time.perf_counter() - started) * 1000)
        log.info(
            "Chunking done %s: %s chunks strategy=%s pages=%s chars min/avg/max=%s/%s/%s",
            relative_path,
            len(chunks),
            strategy,
            pages,
            min(sizes) if sizes else 0,
            int(sum(sizes) / len(sizes)) if sizes else 0,
            max(sizes) if sizes else 0,
            extra={
                "stage": "chunk_ok",
                "suffix": suffix or "none",
                "strategy": strategy,
                "chunks": len(chunks),
                "pages": pages,
                "duration_ms": elapsed,
                "chunk_chars_min": min(sizes) if sizes else 0,
                "chunk_chars_avg": int(sum(sizes) / len(sizes)) if sizes else 0,
                "chunk_chars_max": max(sizes) if sizes else 0,
            },
        )
        return chunks
    finally:
        ingest_source_key.reset(token)


def _chunk_markdown_semantic(file_id: str, fallback: str, text: str, embedder) -> list[Chunk]:
    title = _title_from_text(text, fallback)
    as_of = _as_of_from_text(text)
    chunks: list[Chunk] = []
    idx = 0
    for section, body in split_markdown_sections(text):
        for piece in semantic_section_units(body, embedder):
            if not looks_like_signal(piece):
                continue
            chunks.append(
                make_chunk(
                    file_id=file_id,
                    index=idx,
                    title=title,
                    as_of=as_of,
                    text=piece,
                    strategy="openai_semantic_sd",
                    doc_type="markdown",
                    section=section,
                )
            )
            idx += 1
    return chunks


def _chunk_pdf_semantic(file_id: str, fallback: str, pages: list[tuple[int, str]], embedder) -> list[Chunk]:
    joined = "\n\n".join(text for _, text in pages)
    title = _title_from_text(joined, fallback)
    as_of = _as_of_from_text(joined)
    chunks: list[Chunk] = []
    for i, unit in enumerate(semantic_pdf_units(pages, embedder)):
        chunks.append(
            make_chunk(
                file_id=file_id,
                index=i,
                title=title,
                as_of=as_of,
                text=unit.text,
                strategy="openai_semantic_sd",
                doc_type="pdf",
                page=unit.page,
                section=f"page {unit.page}" if unit.page else "",
            )
        )
    return chunks


def _chunk_markdown(file_id: str, fallback: str, text: str, plan) -> list[Chunk]:
    title = _title_from_text(text, fallback)
    as_of = _as_of_from_text(text)
    chunks: list[Chunk] = []
    for i, (section, body) in enumerate(hybrid_markdown_chunks(text)):
        chunks.append(
            make_chunk(
                file_id=file_id,
                index=i,
                title=title,
                as_of=as_of,
                text=body,
                strategy=plan.strategy,
                doc_type="markdown",
                section=section,
            )
        )
    return chunks


def _chunk_pdf(file_id: str, fallback: str, pages: list[tuple[int, str]], plan) -> list[Chunk]:
    joined = "\n\n".join(text for _, text in pages)
    title = _title_from_text(joined, fallback)
    as_of = _as_of_from_text(joined)
    chunks: list[Chunk] = []
    for i, unit in enumerate(cohesive_pdf_chunks(pages)):
        chunks.append(
            make_chunk(
                file_id=file_id,
                index=i,
                title=title,
                as_of=as_of,
                text=unit.text,
                strategy=plan.strategy,
                doc_type="pdf",
                page=unit.page,
                section=f"page {unit.page}" if unit.page else "",
            )
        )
    return chunks


def _chunk_xlsx(
    file_id: str,
    fallback: str,
    sheets: list[tuple[str, list[str], list[list[str]]]],
    plan,
) -> list[Chunk]:
    chunks: list[Chunk] = []
    idx = 0
    for sheet_name, headers, rows in sheets:
        header_line = "; ".join(h for h in headers if h)
        for start in range(0, len(rows), XLSX_ROWS_PER_CHUNK):
            batch = rows[start : start + XLSX_ROWS_PER_CHUNK]
            lines = [f"# Sheet: {sheet_name}"]
            if header_line:
                lines.append(header_line)
            for row in batch:
                pairs = [f"{h or 'value'}: {v}" for h, v in zip(headers, row) if v]
                if pairs:
                    lines.append("; ".join(pairs))
            text = "\n".join(lines)
            if not looks_like_signal(text) and ":" not in text:
                continue
            chunks.append(
                make_chunk(
                    file_id=file_id,
                    index=idx,
                    title=fallback,
                    as_of="unknown",
                    text=text,
                    strategy=plan.strategy,
                    doc_type="spreadsheet",
                    section=sheet_name,
                )
            )
            idx += 1
    return chunks


def _chunk_jsonl(file_id: str, fallback: str, text: str, plan) -> list[Chunk]:
    as_of = _as_of_from_text(text)
    chunks: list[Chunk] = []
    idx = 0
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        pieces = [line] if len(line) <= 1200 else recursive_split(line)
        for piece in pieces:
            if not looks_like_signal(piece):
                continue
            chunks.append(
                make_chunk(
                    file_id=file_id,
                    index=idx,
                    title=fallback,
                    as_of=as_of,
                    text=piece,
                    strategy=plan.strategy,
                    doc_type="jsonl",
                )
            )
            idx += 1
    return chunks


def _chunk_recursive(file_id: str, fallback: str, text: str, plan, doc_type: str) -> list[Chunk]:
    title = _title_from_text(text, fallback)
    as_of = _as_of_from_text(text)
    chunks: list[Chunk] = []
    for i, piece in enumerate(recursive_split(text) if text else []):
        chunks.append(
            make_chunk(
                file_id=file_id,
                index=i,
                title=title,
                as_of=as_of,
                text=piece,
                strategy=plan.strategy,
                doc_type=doc_type,
            )
        )
    return chunks
